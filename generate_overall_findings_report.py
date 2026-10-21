"""
MADSRCH Case Study — Overall (both genders pooled) findings report generator.

Reads the balanced survey workbook (does NOT modify originals), computes
overall percentages for Male + Female combined (Likert top-box 4–5 % and full
1–5 distributions), exports findings_summary_overall.xlsx, and saves charts
under charts_overall/.

Does NOT import from or modify generate_gender_findings_report.py.

Run:
  python -m pip install -r requirements.txt
  python generate_overall_findings_report.py

Primary input: blo_survey_anonymized.xlsx (SurveyData)
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils.dataframe import dataframe_to_rows

ROOT = Path(__file__).resolve().parent
SOURCE_XLSX = ROOT / "blo_survey_anonymized.xlsx"
if not SOURCE_XLSX.exists():
    SOURCE_XLSX = ROOT / "data" / "survey_data_anonymized.xlsx"
SOURCE_SHEET = "SurveyData"
FIELD_SHEET = "Field_Definitions"
OUT_XLSX = ROOT / "findings_summary_overall.xlsx"
CHART_DIR = ROOT / "charts_overall"

COLOR_OVERALL = "#F4C430"  # saffron
# Likert 1–5 segment colors (disagree/unlikely → agree/likely)
LIKERT_COLORS = {
    1: "#C0392B",
    2: "#E67E22",
    3: "#95A5A6",
    4: "#3498DB",
    5: "#1A5276",
}
LIKERT_SCORES = [1, 2, 3, 4, 5]
AGREE_LEGEND = {
    1: "1 Strongly disagree",
    2: "2 Disagree",
    3: "3 Neutral",
    4: "4 Agree",
    5: "5 Strongly agree",
}
LIKELY_LEGEND = {
    1: "1 Very unlikely",
    2: "2 Unlikely",
    3: "3 Neutral",
    4: "4 Likely",
    5: "5 Very likely",
}
PIE_COLORS = [COLOR_OVERALL, "#5D6D7E"]  # Yes (saffron), Not sure (slate)

LIKERT_TRUST = [
    ("AI_Ads_Creative", "More visually creative\n& appealing"),
    ("AI_Ads_Less_Authentic", "Feel less authentic\n& genuine"),
    ("Trust_AI_Ads_Less", "Trust claims/\ninfo less"),
    ("AI_Ads_Misleading", "More likely to be\nmisleading/deceptive"),
]

LIKERT_ENGAGE = [
    ("Stop_Scrolling", "Stop scrolling\nto look"),
    ("Click_Ad", "Click to learn\nmore"),
    ("Engage_Ad", "Like, comment,\nor share"),
    ("Ignore_Ad", "Ignore the ad\ncompletely"),
]

CRED_LABELS = {
    "It makes me see the brand more positively.": "More positively",
    "It has no effect on my perception of the brand.": "No effect",
    "It makes me see the brand more negatively.": "More negatively",
}
CRED_ORDER = list(CRED_LABELS.keys())

PURCHASE_LABELS = {
    "Much less likely to purchase": "Much less\nlikely",
    "Somewhat less likely to purchase": "Somewhat\nless likely",
    "No effect on my decision": "No effect",
    "Somewhat more likely to purchase": "Somewhat\nmore likely",
    "Much more likely to purchase": "Much more\nlikely",
}
PURCHASE_ORDER = list(PURCHASE_LABELS.keys())

PREF_LABELS = {
    "Definitely the Human-Created ad": "Definitely\nHuman",
    "Somewhat prefer the Human-Created ad": "Somewhat\nHuman",
    "No preference / indifferent": "No\npreference",
    "Somewhat prefer the AI-Made ad": "Somewhat\nAI",
    "Definitely the AI-Made ad": "Definitely\nAI",
}
PREF_ORDER = list(PREF_LABELS.keys())

AWARE_LABELS = {"Yes": "Yes", "I am not sure": "Not sure"}
AWARE_ORDER = ["Yes", "I am not sure"]

SKEPTICISM_LABELS = {
    "Lack of human emotional connection/empathy": "Lack of human\nemotional connection",
    "Fear of data privacy violations (AI using my data to target me).": "Fear of data\nprivacy violations",
    "Concern that the product/image is fake.": "Concern product/\nimage is fake",
    "Belief that the brand is cutting costs/lazy by using AI.": "Brand cutting costs\n/ lazy with AI",
    "I have no concerns, not feeling skeptical.": "No concerns /\nnot skeptical",
}
SKEPTICISM_ORDER = list(SKEPTICISM_LABELS.keys())

TRUST_FACTOR_LABELS = {
    "The reputation of the brand posting the advertisement": "Brand\nreputation",
    "Clear disclosure that the advertisement is developed with AI": "Clear AI\ndisclosure",
    "The presence of real human testimonials or influencers": "Real human\ntestimonials",
    "High-quality visuals and production value": "High-quality\nvisuals",
    "Recommendations from friends or peers": "Friends /\npeers",
}
TRUST_FACTOR_ORDER = list(TRUST_FACTOR_LABELS.keys())


def load_data() -> tuple[pd.DataFrame, pd.DataFrame]:
    if not SOURCE_XLSX.exists():
        raise FileNotFoundError(f"Missing source workbook: {SOURCE_XLSX}")
    df = pd.read_excel(SOURCE_XLSX, sheet_name=SOURCE_SHEET)
    fields = pd.read_excel(SOURCE_XLSX, sheet_name=FIELD_SHEET)
    df = df[df["Gender"].isin(["Male", "Female"])].copy()
    return df, fields


def overall_pct(df: pd.DataFrame, col: str, order: Iterable[str] | None = None) -> pd.Series:
    """% of all respondents choosing each category."""
    n = len(df)
    counts = df[col].value_counts(dropna=False)
    pct = (counts / n * 100).round(1)
    if order is not None:
        ordered = [c for c in order if c in pct.index]
        extra = [c for c in pct.index if c not in ordered]
        pct = pct.reindex(ordered + extra)
    return pct


def overall_counts(df: pd.DataFrame, col: str, order: Iterable[str] | None = None) -> pd.Series:
    counts = df[col].value_counts(dropna=False)
    if order is not None:
        ordered = [c for c in order if c in counts.index]
        extra = [c for c in counts.index if c not in ordered]
        counts = counts.reindex(ordered + extra)
    return counts


def topbox_pct_overall(df: pd.DataFrame, cols: list[str], threshold: int = 4) -> pd.Series:
    """% of all respondents who rated >= threshold (default 4–5)."""
    n = len(df)
    out = {c: (df[c] >= threshold).sum() / n * 100 if n else np.nan for c in cols}
    return pd.Series(out).round(1)


def topbox_counts_overall(df: pd.DataFrame, cols: list[str], threshold: int = 4) -> pd.Series:
    """Headcounts of all respondents who rated >= threshold (default 4–5)."""
    return pd.Series({c: int((df[c] >= threshold).sum()) for c in cols})


def likert_full_pct_overall(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    """% of all respondents at each Likert score 1–5 (rows=items, columns=1..5)."""
    n = len(df)
    rows: dict[str, dict[int, float]] = {}
    for c in cols:
        rows[c] = {}
        for score in LIKERT_SCORES:
            rows[c][score] = (df[c] == score).sum() / n * 100 if n else np.nan
    return pd.DataFrame(rows).T[LIKERT_SCORES].round(1)


def likert_full_counts_overall(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    """Headcounts of all respondents at each Likert score 1–5 (rows=items, columns=1..5)."""
    rows: dict[str, dict[int, int]] = {}
    for c in cols:
        rows[c] = {score: int((df[c] == score).sum()) for score in LIKERT_SCORES}
    return pd.DataFrame(rows).T[LIKERT_SCORES]


def multiselect_pct_overall(
    df: pd.DataFrame,
    col: str,
    options: list[str],
) -> tuple[pd.Series, pd.Series]:
    """
    % and counts of all respondents who selected each option.
    Options are matched as substrings (handles comma-joined multi-select).
    """
    n = len(df)
    texts = df[col].fillna("").astype(str)
    pct_vals: dict[str, float] = {}
    count_vals: dict[str, int] = {}
    for opt in options:
        hits = int(texts.str.contains(opt, regex=False).sum())
        count_vals[opt] = hits
        pct_vals[opt] = hits / n * 100 if n else np.nan
    return pd.Series(pct_vals).round(1), pd.Series(count_vals)


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_df(ws, df: pd.DataFrame, index: bool = True) -> None:
    for r in dataframe_to_rows(df, index=index, header=True):
        ws.append(r)
    style_header(ws)
    ws.freeze_panes = "A2"


def series_to_table(pct: pd.Series, counts: pd.Series | None = None) -> pd.DataFrame:
    """One-row Overall table with categories as columns."""
    row = pct.to_frame().T
    row.index = ["Overall"]
    if counts is not None:
        # returned separately in workbook; this helper is for % rows
        pass
    return row


def styled_single_bars(
    values: pd.Series,
    categories: list[str],
    title: str,
    subtitle: str,
    outfile: Path,
    ylabel: str = "% of Respondents",
) -> None:
    """
    Presentation single bars (Overall), value labels, title + subtitle.
    `values` is a Series keyed by category keys; `categories` are display labels.
    """
    keys = list(values.index)
    if len(categories) != len(keys):
        raise ValueError(
            f"categories ({len(categories)}) must match values ({len(keys)}): {keys}"
        )
    heights = [float(values.iloc[i]) if pd.notna(values.iloc[i]) else 0.0 for i in range(len(keys))]
    x = np.arange(len(keys))
    fig_w = 12.5 if len(keys) >= 5 else 11
    fig, ax = plt.subplots(figsize=(fig_w, 6.2))
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    bars = ax.bar(x, heights, width=0.55, color=COLOR_OVERALL, label="Overall", zorder=3)
    for rect, h in zip(bars, heights):
        ax.annotate(
            f"{h:.1f}%",
            xy=(rect.get_x() + rect.get_width() / 2, h),
            xytext=(0, 12),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=14,
            color="#333333",
            clip_on=False,
        )

    ax.set_xticks(x)
    ax.set_xticklabels(categories, fontsize=10)
    ax.set_ylabel(ylabel, fontsize=11)
    ymax = min(100, max(max(heights) + 18, max(heights) * 1.4 if heights else 10, 10))
    ax.set_ylim(0, ymax)
    ax.yaxis.grid(True, linestyle="-", linewidth=0.6, color="#D0D0D0", zorder=0)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#888888")
    ax.spines["bottom"].set_color("#888888")

    ax.set_title(title, fontsize=14, fontweight="bold", pad=28)
    ax.text(
        0.5,
        1.02,
        subtitle,
        transform=ax.transAxes,
        ha="center",
        va="bottom",
        fontsize=10,
        color="#555555",
        style="italic",
    )
    ax.legend(frameon=False, loc="upper left", bbox_to_anchor=(1.02, 1))
    fig.tight_layout(rect=[0, 0, 0.88, 0.96])
    outfile.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(outfile, dpi=160, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def _pie_label_color(hex_color: str) -> str:
    """White text on dark slices, dark text on light slices."""
    h = hex_color.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    # Relative luminance
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return "#FFFFFF" if luminance < 0.62 else "#1A1A1A"


def _style_pie_pct_labels(
    autotexts,
    texts,
    wedges,
    heights: list[float],
    slice_colors: list[str],
) -> None:
    """Keep % labels inside slices; nudge small-slice % outward (still inside)."""
    for autotext, text, wedge, h, color in zip(
        autotexts, texts, wedges, heights, slice_colors
    ):
        text.set_fontsize(13)
        text.set_color("#333333")
        text.set_fontweight("bold")
        text.set_clip_on(False)
        autotext.set_fontweight("bold")
        autotext.set_clip_on(False)
        autotext.set_ha("center")
        autotext.set_va("center")
        if h < 15:
            autotext.set_fontsize(15)
            ang = np.deg2rad((wedge.theta2 + wedge.theta1) / 2.0)
            r = 0.72
            autotext.set_position((r * np.cos(ang), r * np.sin(ang)))
        else:
            autotext.set_fontsize(16)
        autotext.set_color(_pie_label_color(color))


def styled_pie(
    values: pd.Series,
    labels: list[str],
    title: str,
    subtitle: str,
    outfile: Path,
    colors: list[str] | None = None,
) -> None:
    """Single pie chart with % labels inside slices (no n counts, no explode)."""
    keys = list(values.index)
    if len(labels) != len(keys):
        raise ValueError(f"labels ({len(labels)}) must match values ({len(keys)})")
    heights = [float(values.iloc[i]) if pd.notna(values.iloc[i]) else 0.0 for i in range(len(keys))]
    slice_colors = colors if colors is not None else PIE_COLORS[: len(keys)]

    def autopct_fmt(pct: float) -> str:
        if pct <= 0:
            return ""
        return f"{pct:.1f}%"

    fig, ax = plt.subplots(figsize=(8.8, 6.4))
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    wedges, texts, autotexts = ax.pie(
        heights,
        labels=labels,
        colors=slice_colors,
        autopct=autopct_fmt,
        startangle=70,
        pctdistance=0.55,
        labeldistance=1.12,
        textprops={"fontsize": 12},
        wedgeprops={"linewidth": 2, "edgecolor": "white"},
    )
    _style_pie_pct_labels(autotexts, texts, wedges, heights, slice_colors)
    # Figure-level title/subtitle so they do not overlap (same pattern as gender pies)
    fig.suptitle(title, fontsize=14, fontweight="bold", y=0.97)
    fig.text(0.5, 0.88, subtitle, ha="center", va="top", fontsize=10, color="#555555", style="italic")
    ax.set(aspect="equal")
    # Room inside axes for outside category labels (avoids edge clipping)
    ax.set_xlim(-1.45, 1.45)
    ax.set_ylim(-1.45, 1.45)
    fig.subplots_adjust(left=0.08, right=0.92, bottom=0.08, top=0.74)
    outfile.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(outfile, dpi=160, facecolor="white", bbox_inches="tight", pad_inches=0.55)
    plt.close(fig)


def styled_stacked_likert(
    full_pct: pd.DataFrame,
    item_labels: list[str],
    title: str,
    subtitle: str,
    outfile: Path,
    legend_map: dict[int, str],
) -> None:
    """100% stacked horizontal bars for Likert 1–5 (one bar per item)."""
    items = list(full_pct.index)
    if len(item_labels) != len(items):
        raise ValueError(
            f"item_labels ({len(item_labels)}) must match rows ({len(items)})"
        )
    y = np.arange(len(items))
    fig, ax = plt.subplots(figsize=(11.5, 5.8))
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    left = np.zeros(len(items))
    for score in LIKERT_SCORES:
        widths = [
            float(full_pct.loc[item, score]) if pd.notna(full_pct.loc[item, score]) else 0.0
            for item in items
        ]
        bars = ax.barh(
            y,
            widths,
            left=left,
            height=0.62,
            color=LIKERT_COLORS[score],
            label=legend_map[score],
            zorder=3,
        )
        for rect, w, L in zip(bars, widths, left):
            if w >= 8:
                ax.text(
                    L + w / 2,
                    rect.get_y() + rect.get_height() / 2,
                    f"{w:.0f}%",
                    ha="center",
                    va="center",
                    fontsize=9,
                    color="white" if score in (1, 5) else "#222222",
                )
        left = left + np.array(widths)

    ax.set_yticks(y)
    ax.set_yticklabels(item_labels, fontsize=10)
    ax.set_xlim(0, 100)
    ax.set_xlabel("% of Respondents", fontsize=11)
    ax.xaxis.grid(True, linestyle="-", linewidth=0.6, color="#D0D0D0", zorder=0)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#888888")
    ax.spines["bottom"].set_color("#888888")
    ax.invert_yaxis()
    ax.set_title(title, fontsize=14, fontweight="bold", pad=28)
    ax.text(
        0.5,
        1.02,
        subtitle,
        transform=ax.transAxes,
        ha="center",
        va="bottom",
        fontsize=10,
        color="#555555",
        style="italic",
    )
    ax.legend(frameon=False, loc="upper left", bbox_to_anchor=(1.02, 1), fontsize=9)
    fig.tight_layout(rect=[0, 0, 0.82, 0.96])
    outfile.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(outfile, dpi=160, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def clear_old_overall_charts() -> None:
    if CHART_DIR.exists():
        for p in CHART_DIR.glob("*.png"):
            p.unlink()


def main() -> None:
    df, fields = load_data()
    n_total = len(df)
    n_by_gender = df["Gender"].value_counts().reindex(["Male", "Female"])
    n_dup_total = int((df["Is_Duplicate"] == "Yes").sum())
    n_dup = (
        df.groupby("Gender")["Is_Duplicate"]
        .apply(lambda s: (s == "Yes").sum())
        .reindex(["Male", "Female"])
    )

    # --- Core tables (overall) ---
    awareness_pct = overall_pct(df, "Seen_AI_Labeled_Ads", AWARE_ORDER)
    awareness_n = overall_counts(df, "Seen_AI_Labeled_Ads", AWARE_ORDER)

    credibility_pct = overall_pct(df, "Brand_Credibility_Reaction", CRED_ORDER)
    credibility_n = overall_counts(df, "Brand_Credibility_Reaction", CRED_ORDER)

    trust_cols = [c for c, _ in LIKERT_TRUST]
    trust_topbox = topbox_pct_overall(df, trust_cols)
    trust_topbox_n = topbox_counts_overall(df, trust_cols)
    trust_full = likert_full_pct_overall(df, trust_cols)
    trust_full_n = likert_full_counts_overall(df, trust_cols)
    trust_labels = [lab for _, lab in LIKERT_TRUST]

    skepticism_pct, skepticism_n = multiselect_pct_overall(df, "Skepticism_Concerns", SKEPTICISM_ORDER)

    engage_cols = [c for c, _ in LIKERT_ENGAGE]
    engage_topbox = topbox_pct_overall(df, engage_cols)
    engage_topbox_n = topbox_counts_overall(df, engage_cols)
    engage_full = likert_full_pct_overall(df, engage_cols)
    engage_full_n = likert_full_counts_overall(df, engage_cols)
    engage_labels = [lab for _, lab in LIKERT_ENGAGE]

    pref_pct = overall_pct(df, "Preferred_Ad_Type", PREF_ORDER)
    pref_n = overall_counts(df, "Preferred_Ad_Type", PREF_ORDER)

    purchase_pct = overall_pct(df, "Purchase_Likelihood_AI", PURCHASE_ORDER)
    purchase_n = overall_counts(df, "Purchase_Likelihood_AI", PURCHASE_ORDER)

    trust_factor_pct = overall_pct(df, "Top_Trust_Factor", TRUST_FACTOR_ORDER)
    trust_factor_n = overall_counts(df, "Top_Trust_Factor", TRUST_FACTOR_ORDER)

    neg_cred = float(credibility_pct.get("It makes me see the brand more negatively.", 0.0))
    less_buy = float(
        purchase_pct.get("Much less likely to purchase", 0.0)
        + purchase_pct.get("Somewhat less likely to purchase", 0.0)
    )
    prefer_human = float(
        pref_pct.get("Definitely the Human-Created ad", 0.0)
        + pref_pct.get("Somewhat prefer the Human-Created ad", 0.0)
    )

    # --- Charts (charts_overall/ only) ---
    clear_old_overall_charts()
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    chart_specs: list[tuple[str, str, str]] = []

    p = CHART_DIR / "00_awareness.png"
    aware_keys = [c for c in AWARE_ORDER if c in awareness_pct.index]
    aware_vals = awareness_pct[aware_keys]
    styled_pie(
        aware_vals,
        [AWARE_LABELS[c] for c in aware_keys],
        "Awareness of AI-Labeled Ads",
        "% who said Yes vs Not sure (valid sample, both genders)",
        p,
    )
    chart_specs.append(("Awareness (pie)", "Awareness", p.name))

    p = CHART_DIR / "01_trust_credibility_topbox.png"
    styled_single_bars(
        trust_topbox,
        trust_labels,
        "Trust & Credibility Toward AI-Labeled Ads",
        "% who Agree/Strongly Agree (rating 4-5 out of 5)",
        p,
    )
    chart_specs.append(("Trust & credibility (Agree 4–5)", "Trust_Credibility", p.name))

    p = CHART_DIR / "01b_trust_credibility_full.png"
    styled_stacked_likert(
        trust_full,
        trust_labels,
        "Trust & Credibility Toward AI-Labeled Ads (Full Scale)",
        "% at each rating 1–5 (Strongly disagree → Strongly agree)",
        p,
        AGREE_LEGEND,
    )
    chart_specs.append(("Trust & credibility (full 1–5)", "Trust_Credibility", p.name))

    p = CHART_DIR / "02_brand_reaction.png"
    cred_keys = [c for c in CRED_ORDER if c in credibility_pct.index]
    styled_single_bars(
        credibility_pct[cred_keys],
        [CRED_LABELS[c] for c in cred_keys],
        "Immediate Brand Reaction to an AI Label",
        "% of respondents choosing each reaction",
        p,
    )
    chart_specs.append(("Brand reaction", "Trust_Credibility", p.name))

    p = CHART_DIR / "03_skepticism_concerns.png"
    skep_keys = list(SKEPTICISM_ORDER)
    styled_single_bars(
        skepticism_pct[skep_keys],
        [SKEPTICISM_LABELS[c] for c in skep_keys],
        "Concerns Driving Skepticism Toward AI-Labeled Ads",
        "% selecting each concern (multi-select; overall sample)",
        p,
    )
    chart_specs.append(("Skepticism concerns", "Skepticism", p.name))

    p = CHART_DIR / "04_engagement_topbox.png"
    styled_single_bars(
        engage_topbox,
        engage_labels,
        "Engagement With AI-Labeled Ads",
        "% who rated Likely/Very likely (4-5 out of 5)",
        p,
    )
    chart_specs.append(("Engagement (Likely 4–5)", "Engagement", p.name))

    p = CHART_DIR / "04b_engagement_full.png"
    styled_stacked_likert(
        engage_full,
        engage_labels,
        "Engagement With AI-Labeled Ads (Full Scale)",
        "% at each rating 1–5 (Very unlikely → Very likely)",
        p,
        LIKELY_LEGEND,
    )
    chart_specs.append(("Engagement (full 1–5)", "Engagement", p.name))

    p = CHART_DIR / "05_preferred_ad_type.png"
    pref_keys = [c for c in PREF_ORDER if c in pref_pct.index]
    styled_single_bars(
        pref_pct[pref_keys],
        [PREF_LABELS[c] for c in pref_keys],
        "Which Ad Would Students Engage With?",
        "% choosing Human-created vs AI-made (same product)",
        p,
    )
    chart_specs.append(("Preferred ad type", "Engagement", p.name))

    p = CHART_DIR / "06_purchase_likelihood.png"
    purch_keys = [c for c in PURCHASE_ORDER if c in purchase_pct.index]
    styled_single_bars(
        purchase_pct[purch_keys],
        [PURCHASE_LABELS[c] for c in purch_keys],
        "Purchase Likelihood if the Ad Used AI",
        "% of respondents by purchase effect",
        p,
    )
    chart_specs.append(("Purchase likelihood", "Purchase", p.name))

    p = CHART_DIR / "07_top_trust_factor.png"
    factor_keys = [c for c in TRUST_FACTOR_ORDER if c in trust_factor_pct.index]
    styled_single_bars(
        trust_factor_pct[factor_keys],
        [TRUST_FACTOR_LABELS[c] for c in factor_keys],
        "Most Important Factor for Trusting an Online Ad",
        "% choosing each factor (select only one)",
        p,
    )
    chart_specs.append(("Top trust factor", "Top_Trust_Factor", p.name))

    # --- Workbook ---
    wb = Workbook()

    ws = wb.active
    ws.title = "Column_Map"
    cmap = fields.rename(columns={"Tableau_Field": "Field", "Original_Question": "Question"})
    write_df(ws, cmap, index=False)

    ws = wb.create_sheet("Sample_Profile")
    ws.append(["Source", SOURCE_XLSX.name])
    ws.append(["Sheet", SOURCE_SHEET])
    ws.append(["n (overall)", int(n_total)])
    ws.append(["Male", int(n_by_gender["Male"])])
    ws.append(["Female", int(n_by_gender["Female"])])
    ws.append(["Duplicates overall (Is_Duplicate=Yes)", n_dup_total])
    ws.append(["Male duplicates (Is_Duplicate=Yes)", int(n_dup["Male"])])
    ws.append(["Female duplicates (Is_Duplicate=Yes)", int(n_dup["Female"])])
    style_header(ws)

    def append_series_block(ws, title: str, series: pd.Series, index_name: str = "Category") -> None:
        ws.append([title])
        table = series.rename("Overall").reset_index()
        table.columns = [index_name, "Overall"]
        for r in dataframe_to_rows(table, index=False, header=True):
            ws.append(r)

    ws = wb.create_sheet("Awareness")
    append_series_block(ws, "Counts", awareness_n)
    ws.append([])
    append_series_block(ws, "% of all respondents", awareness_pct)
    style_header(ws)

    ws = wb.create_sheet("Trust_Credibility")
    append_series_block(ws, "Brand reaction %", credibility_pct)
    ws.append([])
    append_series_block(ws, "Brand reaction counts", credibility_n)
    ws.append([])
    append_series_block(ws, "Likert top-box % (rated 4–5)", trust_topbox, index_name="Item")
    ws.append([])
    append_series_block(ws, "Likert top-box counts (rated 4–5)", trust_topbox_n, index_name="Item")
    ws.append([])
    ws.append(["Likert full scale % (ratings 1–5)"])
    full_trust = trust_full.copy()
    full_trust.columns = [f"{s} ({AGREE_LEGEND[s].split(' ', 1)[1]})" for s in LIKERT_SCORES]
    full_trust = full_trust.reset_index().rename(columns={"index": "Item"})
    for r in dataframe_to_rows(full_trust, index=False, header=True):
        ws.append(r)
    ws.append([])
    ws.append(["Likert full scale counts (ratings 1–5)"])
    full_trust_n = trust_full_n.copy()
    full_trust_n.columns = [f"{s} ({AGREE_LEGEND[s].split(' ', 1)[1]})" for s in LIKERT_SCORES]
    full_trust_n = full_trust_n.reset_index().rename(columns={"index": "Item"})
    for r in dataframe_to_rows(full_trust_n, index=False, header=True):
        ws.append(r)
    style_header(ws)

    ws = wb.create_sheet("Skepticism")
    append_series_block(ws, "Skepticism concern counts (multi-select)", skepticism_n)
    ws.append([])
    append_series_block(
        ws,
        "Skepticism concern % of all respondents (multi-select)",
        skepticism_pct,
    )
    style_header(ws)

    ws = wb.create_sheet("Engagement")
    append_series_block(ws, "Engagement top-box % (rated 4–5)", engage_topbox, index_name="Item")
    ws.append([])
    append_series_block(ws, "Engagement top-box counts (rated 4–5)", engage_topbox_n, index_name="Item")
    ws.append([])
    ws.append(["Engagement full scale % (ratings 1–5)"])
    full_engage = engage_full.copy()
    full_engage.columns = [f"{s} ({LIKELY_LEGEND[s].split(' ', 1)[1]})" for s in LIKERT_SCORES]
    full_engage = full_engage.reset_index().rename(columns={"index": "Item"})
    for r in dataframe_to_rows(full_engage, index=False, header=True):
        ws.append(r)
    ws.append([])
    ws.append(["Engagement full scale counts (ratings 1–5)"])
    full_engage_n = engage_full_n.copy()
    full_engage_n.columns = [f"{s} ({LIKELY_LEGEND[s].split(' ', 1)[1]})" for s in LIKERT_SCORES]
    full_engage_n = full_engage_n.reset_index().rename(columns={"index": "Item"})
    for r in dataframe_to_rows(full_engage_n, index=False, header=True):
        ws.append(r)
    ws.append([])
    append_series_block(ws, "Preferred ad type %", pref_pct)
    ws.append([])
    append_series_block(ws, "Preferred ad type counts", pref_n)
    style_header(ws)

    ws = wb.create_sheet("Purchase")
    append_series_block(ws, "Purchase likelihood %", purchase_pct)
    ws.append([])
    append_series_block(ws, "Purchase likelihood counts", purchase_n)
    style_header(ws)

    ws = wb.create_sheet("Top_Trust_Factor")
    append_series_block(ws, "Top trust factor %", trust_factor_pct)
    ws.append([])
    append_series_block(ws, "Top trust factor counts", trust_factor_n)
    style_header(ws)

    ws = wb.create_sheet("Chart_Index")
    ws.append(["Chart", "Sheet", "PNG"])
    for title, sheet, png in chart_specs:
        ws.append([title, sheet, f"charts_overall/{png}"])
    style_header(ws)

    ws = wb.create_sheet("Limitations")
    for note in [
        "Convenience sample of DLSU undergraduates; descriptive % only.",
        "Pooled Male + Female (n=184); 92/92 balance includes duplicated rows (see Sample_Profile).",
        "Valid sample excludes people who said they never saw AI-labeled ads.",
        "Likert charts include top-box % (rated 4–5) and full 1–5 distribution charts/tables.",
        "Skepticism_Concerns is multi-select; % can sum over 100% overall.",
        "Gender comparisons are in the separate Male vs Female report (findings_summary.xlsx).",
        "Original survey Excels were not modified.",
    ]:
        ws.append([note])

    wb.save(OUT_XLSX)

    print("=== MADSRCH overall findings report ===")
    print(f"Source: {SOURCE_XLSX.name} | n={n_total} (Male+Female pooled)")
    print(f"Wrote: {OUT_XLSX.name}")
    print(f"Charts: {len(chart_specs)} files in charts_overall/")
    for _, _, png in chart_specs:
        print(f"  - {png}")
    print("\nTrust top-box % (4–5):")
    print(trust_topbox.to_string())
    print("\nTrust full 1–5 %:")
    print(trust_full.to_string())
    print("\nEngagement top-box % (4–5):")
    print(engage_topbox.to_string())
    print("\nEngagement full 1–5 %:")
    print(engage_full.to_string())
    print("\nSkepticism concern %:")
    print(skepticism_pct.to_string())
    print("\nTop trust factor %:")
    print(trust_factor_pct.to_string())
    print("\nAwareness %:")
    print(awareness_pct.to_string())
    print("\nBrand reaction %:")
    print(credibility_pct.to_string())
    print("\nPurchase %:")
    print(purchase_pct.to_string())
    print("\nPreferred ad %:")
    print(pref_pct.to_string())
    print("\nCollapsed summaries:")
    print("negative reaction", round(neg_cred, 1))
    print("prefer human", round(prefer_human, 1))
    print("less likely buy", round(less_buy, 1))
    print("\nDone.")


if __name__ == "__main__":
    main()
